"""Importa al motor las respuestas del cuestionario de reglas.

Uso (desde la raíz del repo):
    python herramientas/importar_respuestas.py --nube        # lo último que respondió negocio en línea
    python herramientas/importar_respuestas.py --nube --ver  # solo muestra quién respondió y cuánto, sin escribir
    python herramientas/importar_respuestas.py ruta/al/excel-de-respuestas.xlsx

Con --nube lee la tabla "respuestas" de Azure Storage (cuenta motorcarterarespuestas), donde el
cuestionario publicado guarda cada avance; toma el guardado más reciente. Necesita la CLI de Azure
con sesión iniciada (az login). Con un Excel lee su hoja "Respuestas". En ambos casos escribe:
  - config/reglas-negocio.properties: los ajustes que el motor ya sabe aplicar
    (ReglasNegocio.CLAVES). ReglasNegocioConfiguradaTest corre los 9 casos con ellos.
  - docs/respuestas-negocio.md: todas las respuestas, con lo que queda por desarrollar.
Después: ./mvnw -q -B test y revisar las pruebas que cambien frente al Excel.
"""
import json
import shutil
import subprocess
import sys
import urllib.request
from datetime import datetime, timedelta, timezone
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
# Debe coincidir con ReglasNegocio.CLAVES.
CLAVES_MOTOR = {'prelacion', 'formulaMora', 'baseInteres', 'topeCobranzaPorPago', 'baseServicios', 'serviciosDespuesDelPlazo', 'moraAntesDeRecuperacion',
                'aplicacionAbono', 'interesPrepagoPorDias', 'margenUsura', 'modalidadUsura'}
# Preguntas cuyos ajustes NO se aplican porque el comentario de negocio contradice la opción marcada
# (la opción era de relleno: el formulario no guardaba sin elegir una). Se versiona aquí; para
# descartar otra pregunta basta agregar su ID. Aparecen en el md como "Opción contradicha por el comentario".
DESCARTADOS = {
    'A3': 'abonoMinimo=0: Diego no dio un mínimo; el abono sigue la prelación de Inputs (pregunta abierta)',
    'R2': 'costoRecuperacion=0: los costos de recuperación van dentro de los gastos de cobranza, aún sin definir',
    'T5': 'saldo-total: el comentario pide investigar las cláusulas aceleratorias; se repitió sin responder (ronda 4)',
    'C5': 'valor-total: el detalle (total = moto + alistamiento) contradice la opción con GPS y gastos comerciales; se repitió sin responder (ronda 4)',
}
# Código de la opción que marca una pregunta aún sin opción real; sus ajustes tampoco se aplican.
CODIGO_PENDIENTE = 'PENDIENTE'
ENCABEZADOS = ['ID', 'Bloque', 'Pregunta', 'Código', 'Respuesta', 'Valor', 'Unidad', 'Detalle', 'Comentario',
               'Ajustes al motor', 'Completa']


CUENTA_NUBE = 'motorcarterarespuestas'
TABLA_NUBE = 'respuestas'
GRUPO_NUBE = 'rg-motor-cartera'


def _az(*args):
    az = shutil.which('az') or shutil.which('az.cmd')
    if not az:
        sys.exit('Falta la CLI de Azure (az) con sesión iniciada: az login')
    salida = subprocess.run([az, *args], capture_output=True)
    if salida.returncode != 0:
        sys.exit(f'Azure CLI falló: {salida.stderr.decode("utf-8", "replace").strip()}')
    return salida.stdout.decode('ascii').strip()


def leer_nube():
    """Guardados del cuestionario en línea, del más reciente al más antiguo (sin las pruebas técnicas).

    La CLI de Azure solo firma un permiso de lectura de una hora; la tabla se lee por HTTPS, donde las
    respuestas llegan en UTF-8 sin pasar por la consola de Windows.
    """
    vence = (datetime.now(timezone.utc) + timedelta(hours=1)).strftime('%Y-%m-%dT%H:%MZ')
    clave = _az('storage', 'account', 'keys', 'list', '--account-name', CUENTA_NUBE, '--resource-group', GRUPO_NUBE,
                '--query', '[0].value', '-o', 'tsv')
    permiso = _az('storage', 'table', 'generate-sas', '--name', TABLA_NUBE, '--account-name', CUENTA_NUBE,
                  '--account-key', clave, '--permission', 'r', '--expiry', vence, '--https-only', '-o', 'tsv')
    base = f'https://{CUENTA_NUBE}.table.core.windows.net/{TABLA_NUBE}()?{permiso}'
    guardados, continuacion = [], ''
    while True:
        solicitud = urllib.request.Request(base + continuacion, headers={
            'Accept': 'application/json;odata=nometadata', 'x-ms-version': '2019-02-02'})
        with urllib.request.urlopen(solicitud) as respuesta:
            guardados += json.loads(respuesta.read().decode('utf-8'))['value']
            particion = respuesta.headers.get('x-ms-continuation-NextPartitionKey')
            fila = respuesta.headers.get('x-ms-continuation-NextRowKey')
        if not particion:
            break
        continuacion = f'&NextPartitionKey={particion}&NextRowKey={fila or ""}'
    guardados = [e for e in guardados
                 if e.get('PartitionKey') == 'respuestas' and not str(e.get('sesion', '')).startswith('prueba-')]
    return sorted(guardados, key=lambda e: int(str(e['RowKey']).split('-', 1)[0]), reverse=True)


def filas_de_guardado(guardado):
    encabezados = json.loads(guardado['encabezados'])
    if encabezados[:len(ENCABEZADOS)] != ENCABEZADOS:
        sys.exit(f'Encabezados inesperados en la nube: {encabezados}')
    filas = [dict(zip(ENCABEZADOS, ['' if v is None else v for v in fila])) for fila in json.loads(guardado['filas'])]
    fecha = datetime.fromtimestamp(int(str(guardado['RowKey']).split('-', 1)[0]) / 1000)
    resumen = {'Respondió': guardado.get('nombre', ''), 'Cargo': guardado.get('cargo', ''),
               'Fecha': f'{fecha:%Y-%m-%d %H:%M}' + (' (enviado)' if guardado.get('final') else ' (en progreso)')}
    return filas, resumen


def leer(ruta):
    import openpyxl

    libro = openpyxl.load_workbook(ruta, data_only=True)
    if 'Respuestas' not in libro.sheetnames:
        sys.exit('El archivo no tiene la hoja "Respuestas": ¿es el Excel que genera el cuestionario?')
    hoja = libro['Respuestas']
    encabezados = [c.value for c in hoja[1]]
    if encabezados[:len(ENCABEZADOS)] != ENCABEZADOS:
        sys.exit(f'Encabezados inesperados en "Respuestas": {encabezados}')
    filas = []
    for fila in hoja.iter_rows(min_row=2, values_only=True):
        if fila[0]:
            filas.append(dict(zip(ENCABEZADOS, ['' if v is None else v for v in fila])))
    resumen = {}
    if 'Resumen' in libro.sheetnames:
        for a, b in libro['Resumen'].iter_rows(min_row=1, max_col=2, values_only=True):
            if a:
                resumen[str(a)] = '' if b is None else str(b)
    return filas, resumen


def ajustes(fila):
    resultado = {}
    for parte in str(fila['Ajustes al motor']).split(';'):
        if '=' in parte:
            clave, valor = parte.split('=', 1)
            resultado[clave.strip()] = valor.strip()
    return resultado


def descartada(fila):
    """Motivo por el que los ajustes de la fila no se aplican, o None si sí se aplican."""
    if fila['ID'] in DESCARTADOS:
        return DESCARTADOS[fila['ID']]
    codigo = str(fila['Código']).strip()
    if not codigo:
        return 'fila sin opción'
    if codigo.upper() == CODIGO_PENDIENTE:
        return 'opción PENDIENTE'
    return None


def config_actual():
    """Reglas ya configuradas: las respuestas nuevas se suman a ellas, no las borran."""
    archivo = RAIZ / 'config' / 'reglas-negocio.properties'
    actuales = {}
    if archivo.exists():
        for linea in archivo.read_text(encoding='utf-8').splitlines():
            if '=' in linea and not linea.lstrip().startswith('#'):
                clave, valor = linea.split('=', 1)
                actuales[clave.strip()] = (valor.strip(), 'anterior')
    return actuales


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
    args = sys.argv[1:]
    if not args:
        sys.exit(__doc__)
    if args[0] == '--nube':
        guardados = leer_nube()
        if not guardados:
            print('Todavía no hay respuestas en la nube.')
            return
        ultimo = guardados[0]
        filas, resumen = filas_de_guardado(ultimo)
        print(f'{len(guardados)} guardados en la nube. Último: {resumen["Respondió"]} ({resumen["Cargo"]}) · '
              f'{resumen["Fecha"]} · {ultimo.get("completas")} de {ultimo.get("total")} respondidas.')
        if '--ver' in args:
            return
        fuente = f'la nube ({resumen["Fecha"]})'
    else:
        ruta = Path(args[0])
        filas, resumen = leer(ruta)
        fuente = ruta.name

    aplicables, por_desarrollar, contradichas = config_actual(), {}, []
    for f in filas:
        motivo = descartada(f)
        if motivo:
            if f['ID'] in DESCARTADOS:
                contradichas.append((f['ID'], motivo))
            continue
        for clave, valor in ajustes(f).items():
            (aplicables if clave in CLAVES_MOTOR else por_desarrollar)[clave] = (valor, f['ID'])

    (RAIZ / 'config').mkdir(exist_ok=True)
    props = [f'# Generado desde {fuente} el {datetime.now():%Y-%m-%d %H:%M}',
             f'# Respondió: {resumen.get("Respondió", "")} ({resumen.get("Cargo", "")})']
    props += [f'{clave}={valor}' for clave, (valor, _) in sorted(aplicables.items())]
    (RAIZ / 'config' / 'reglas-negocio.properties').write_text('\n'.join(props) + '\n', encoding='utf-8')

    completas = sum(1 for f in filas if f['Completa'] == 'Sí')
    md = [
        '# Respuestas de negocio al cuestionario de reglas',
        '',
        f'Fuente: {fuente} · Respondió: {resumen.get("Respondió", "")} ({resumen.get("Cargo", "")}) · '
        f'{resumen.get("Fecha", "")} · {completas} de {len(filas)} completas.',
        '',
        '## Aplicado al motor (`config/reglas-negocio.properties`)',
        '',
    ]
    md += [f'- `{c}={v}` (pregunta {p})' for c, (v, p) in sorted(aplicables.items())] or ['- Ninguno.']
    md += ['', '## Decidido pero pendiente de desarrollo en el motor', '']
    md += [f'- `{c}={v}` (pregunta {p})' for c, (v, p) in sorted(por_desarrollar.items())] or ['- Ninguno.']
    md += ['', '## Opción contradicha por el comentario: no se aplica', '']
    md += [f'- {p}: {m}' for p, m in contradichas] or ['- Ninguna.']
    md += ['', '## Todas las respuestas', '', '| ID | Pregunta | Respuesta | Valor | Detalle | Comentario |', '|---|---|---|---|---|---|']
    for f in filas:
        valor = f'{f["Valor"]} {f["Unidad"]}'.strip()
        celdas = [f['ID'], f['Pregunta'], f['Respuesta'] or '*sin responder*', valor, f['Detalle'], f['Comentario']]
        md.append('| ' + ' | '.join(str(c).replace('|', '/').replace('\n', ' ') for c in celdas) + ' |')
    (RAIZ / 'docs' / 'respuestas-negocio.md').write_text('\n'.join(md) + '\n', encoding='utf-8')

    print(f'{completas} de {len(filas)} preguntas completas')
    print(f'Aplicado al motor: {", ".join(sorted(aplicables)) or "nada"}')
    print(f'Pendiente de desarrollo: {", ".join(sorted(por_desarrollar)) or "nada"}')
    print('Escrito config/reglas-negocio.properties y docs/respuestas-negocio.md. Siguiente: ./mvnw -q -B test')


if __name__ == '__main__':
    main()
