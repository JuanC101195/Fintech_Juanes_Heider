"""Llena la columna "Resultado software" de la hoja Validaciones con lo que calcula el motor (HU-17b).

Uso (desde la raíz del repo, con el API corriendo):
    python -I herramientas/llenar_validaciones.py                      # API local http://localhost:8080
    python -I herramientas/llenar_validaciones.py --api                # API desplegado en Azure
    python -I herramientas/llenar_validaciones.py --url http://localhost:8091/api/v1 --salida C:/tmp

Para cada prueba de la tabla principal de Validaciones:
  - F "Resultado software": el valor del motor, leído de GET /api/v1/casos/{n}?perfilReglas=EXCEL.
  - K "Valor Excel": fórmula a la celda de la hoja del caso que tiene el valor esperado.
  - L "Diferencia": |F − K| (en "Exacto", 0 si coinciden y 1 si no).
  - M "Resultado": OK si la diferencia está dentro de la tolerancia (columna G), FALLA si no.
  - N "Origen": de dónde sale cada lado (campo del API y celda del Excel).
Una prueba sin mapeo seguro al motor queda con F vacía, M = "SIN MAPEO" y el motivo en J (Observaciones).

Guarda un archivo NUEVO (nunca sobrescribe el original) con el resto del libro intacto: fórmulas, formato
y validaciones. openpyxl no guarda los valores calculados de las fórmulas, así que el libro se marca para
recalcular al abrirlo en Excel. Solo usa la biblioteca estándar más openpyxl.
"""
import argparse
import json
import sys
import tempfile
import unicodedata
import urllib.request
from copy import copy
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Callable, Optional

import openpyxl
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parent.parent
EXCEL = RAIZ / 'docs' / 'fuentes' / 'casos-operaciones-v2.xlsx'
URL_LOCAL = 'http://localhost:8080/api/v1'
URL_AZURE = 'https://motor-cartera-api-jc.azurewebsites.net/api/v1'
HOJA = 'Validaciones'
FILA_ENCABEZADO = 2
PRIMERA_FILA = 3
# Columnas de la hoja Validaciones (las de Diego van de A a J; K a N las agrega esta herramienta).
COL = {'numero': 1, 'caso': 2, 'variable': 4, 'tipo': 5, 'resultado': 6, 'tolerancia': 7, 'periodo': 8,
       'observaciones': 10, 'esperado': 11, 'diferencia': 12, 'estado': 13, 'origen': 14}
ENCABEZADOS_NUEVOS = {'esperado': 'Valor Excel', 'diferencia': 'Diferencia', 'estado': 'Resultado (OK / FALLA)',
                      'origen': 'Origen (motor ↔ Excel)'}
SIN_MAPEO = 'SIN MAPEO'
TOLERANCIA_POR_TIPO = {'Dinero': Decimal(1), 'Tasa': Decimal('0.000001'), 'Conteo': Decimal(0), 'Exacto': Decimal(0)}


# ----------------------------------------------------------------------------------------------- motor

def movimiento(respuesta: dict, periodo: int) -> dict:
    """El movimiento del periodo pedido (el periodo 0 es el desembolso)."""
    for m in respuesta['movimientos']:
        if m['periodo'] == periodo:
            return m
    raise KeyError(f'el motor no trae el periodo {periodo}')


def campo(ruta: str) -> Callable[[dict, Optional[int]], Decimal]:
    """Un campo del movimiento del periodo; `ruta` admite objetos anidados ("pagoTerceros.fondoGarantias")."""
    def leer(respuesta, periodo):
        valor = movimiento(respuesta, periodo)
        for parte in ruta.split('.'):
            valor = valor[parte]
        return Decimal(str(valor))
    leer.descripcion = f'movimientos[periodo].{ruta}'
    return leer


def _con_descripcion(descripcion):
    def poner(f):
        f.descripcion = descripcion
        return f
    return poner


@_con_descripcion('primer periodo cuya tasaPeriodica difiere de la del periodo 1')
def primer_cambio_de_tasa(respuesta, _periodo=None):
    inicial = Decimal(movimiento(respuesta, 1)['tasaPeriodica'])
    for m in respuesta['movimientos']:
        if m['periodo'] > 1 and abs(Decimal(m['tasaPeriodica']) - inicial) > Decimal('1e-12'):
            return Decimal(m['periodo'])
    raise KeyError('el motor no cambia la tasa en ningún periodo')


@_con_descripcion('periodo del último movimiento')
def ultimo_periodo(respuesta, _periodo=None):
    return Decimal(respuesta['movimientos'][-1]['periodo'])


@_con_descripcion('credito.cuotas − periodo del último movimiento')
def cuotas_eliminadas(respuesta, _periodo=None):
    return Decimal(respuesta['credito']['cuotas']) - ultimo_periodo(respuesta)


@_con_descripcion('movimientos con periodo > credito.cuotas')
def periodos_adicionales(respuesta, _periodo=None):
    cuotas = respuesta['credito']['cuotas']
    return Decimal(sum(1 for m in respuesta['movimientos'] if m['periodo'] > cuotas))


@_con_descripcion('estado del último movimiento')
def estado_final(respuesta, _periodo=None):
    return respuesta['movimientos'][-1]['estado']


def normalizar_estado(texto) -> str:
    """'Vigente al día' y 'VIGENTE_AL_DIA' quedan iguales: sin tildes, mayúsculas y _ por espacios."""
    sin_tildes = unicodedata.normalize('NFKD', str(texto)).encode('ascii', 'ignore').decode()
    return sin_tildes.strip().upper().replace(' ', '_')


# ----------------------------------------------------------------------------------------------- Excel

@dataclass(frozen=True)
class Esperado:
    """Valor esperado: la fórmula que se escribe en la columna K y cómo calcularlo con el libro en valores."""
    formula: str
    valor: Callable[[object], object]


def fila_del_periodo(hoja, periodo: int) -> int:
    for fila in range(PRIMERA_FILA, hoja.max_row + 1):
        if hoja.cell(fila, 1).value == periodo:
            return fila
    raise KeyError(f'{hoja.title}: no hay fila del periodo {periodo}')


def columna_por_encabezado(hoja, encabezado: str) -> int:
    c = 1
    while hoja.cell(FILA_ENCABEZADO, c).value is not None:
        if hoja.cell(FILA_ENCABEZADO, c).value == encabezado:
            return c
        c += 1
    raise KeyError(f'{hoja.title}: no hay columna "{encabezado}"')


def en_columna(nombre_hoja: str, encabezado: str, periodo: int, libro) -> Esperado:
    hoja = libro[nombre_hoja]
    ref = f'{get_column_letter(columna_por_encabezado(hoja, encabezado))}{fila_del_periodo(hoja, periodo)}'
    return Esperado(f'={nombre_hoja}!{ref}', lambda lv, h=nombre_hoja, r=ref: lv[h][r].value)


def en_celda(nombre_hoja: str, ref: str) -> Esperado:
    return Esperado(f'={nombre_hoja}!{ref}', lambda lv: lv[nombre_hoja][ref].value)


@dataclass(frozen=True)
class Mapeo:
    hoja: str                                   # hoja del caso en el Excel
    motor: Callable                             # (respuesta, periodo) → valor del motor
    esperado: Callable                          # (libro, periodo) → Esperado
    nota: str = ''


def columna(encabezado: str, ruta_motor: str, nota: str = '') -> Callable[[str], Mapeo]:
    """Variable que está en la tabla del caso, en la fila del periodo de la prueba."""
    return lambda hoja: Mapeo(hoja, campo(ruta_motor), lambda libro, p: en_columna(hoja, encabezado, p, libro), nota)


def celda(ref: str, motor: Callable, nota: str = '', hoja_ref: Optional[str] = None) -> Callable[[str], Mapeo]:
    """Variable que está en una celda fija (panel de indicadores del caso o Inputs)."""
    return lambda hoja: Mapeo(hoja, motor, lambda libro, p: en_celda(hoja_ref or hoja, ref), nota)


def formula(texto: str, calcular: Callable, motor: Callable, nota: str = '') -> Callable[[str], Mapeo]:
    """Variable sin celda propia en el Excel: se arma con una fórmula sobre celdas del libro."""
    return lambda hoja: Mapeo(hoja, motor, lambda libro, p: Esperado(texto, calcular), nota)


HOJAS = {1: 'Caso1_Normal', 2: 'Caso2_CambioTasa', 3: 'Caso3_Dacion', 4: 'Caso4_Retoma', 5: 'Caso5_Prepago',
         6: 'Caso6_MenorPlazo', 7: 'Caso7_MenorCuota', 8: 'Caso8_PagoInferior', 9: 'Caso9_Cobranza',
         10: 'Caso10_TasaCuotaFija', 11: 'Caso11_PagoInfDacion', 12: 'Caso12_PagoInfNormaliza'}

# Variables que se repiten entre casos: encabezado en la hoja del caso → campo del movimiento.
SALDO = 'Saldo final'
SALDO_FINAL = columna(SALDO, 'saldoFinal')
TASA_PERIODICA = columna('Tasa periódica', 'tasaPeriodica')
CUOTA_FINANCIERA = columna('Cuota financiera', 'cuotaFinanciera')
CUOTAS_VENCIDAS = columna('Cuotas vencidas', 'cuotasVencidas')
DEUDA_A_EXTINGUIR = columna('Deuda a extinguir', 'deudaAExtinguir')
INVENTARIO = columna('Inventario', 'inventario')
CXC_FCC = columna('CxC FCC', 'cxcFondoGarantias')
ABONO_EXTRA = columna('Abono extra', 'abonoExtra')
PAGO_RECIBIDO = columna('Pago recibido (cuota total)', 'pagoRecibido')
CAPITAL_VENCIDO = columna('Capital vencido', 'capitalVencido')
INTERES_MORA = columna('Interés mora', 'interesMora')

# Número de la prueba → (caso, mapeo). El periodo sale de la columna H de la prueba.
MAPEO = {
    1: (1, columna(SALDO, 'saldoFinal', 'periodo 0 = desembolso')),
    2: (1, TASA_PERIODICA),
    3: (1, CUOTA_FINANCIERA),
    4: (1, SALDO_FINAL),
    5: (1, columna('Pago FCC', 'pagoTerceros.fondoGarantias', 'giro mensual al fondo de garantías')),
    6: (2, celda('C43', primer_cambio_de_tasa, hoja_ref='Inputs')),
    7: (2, TASA_PERIODICA),
    8: (2, CUOTA_FINANCIERA),
    9: (2, SALDO_FINAL),
    10: (3, CUOTAS_VENCIDAS),
    11: (3, DEUDA_A_EXTINGUIR),
    12: (3, INVENTARIO),
    13: (3, CXC_FCC),
    14: (4, DEUDA_A_EXTINGUIR),
    15: (4, INVENTARIO),
    16: (4, columna('CxP deudor', 'cxpDeudor')),
    17: (5, columna('Prepago capital', 'prepagoCapital')),
    18: (5, SALDO_FINAL),
    19: (6, ABONO_EXTRA),
    20: (6, celda('V6', ultimo_periodo)),
    21: (6, formula('=Inputs!B19-Caso6_MenorPlazo!V6',
                    lambda lv: lv['Inputs']['B19'].value - lv['Caso6_MenorPlazo']['V6'].value, cuotas_eliminadas,
                    'el Excel no la calcula: cuotas pactadas (Inputs B19) − último periodo con pago (V6)')),
    22: (7, ABONO_EXTRA),
    23: (7, CUOTA_FINANCIERA),
    24: (7, SALDO_FINAL),
    25: (8, PAGO_RECIBIDO),
    26: (8, CAPITAL_VENCIDO),
    27: (8, INTERES_MORA),
    28: (8, formula('=MAX(0,Caso8_PagoInferior!AI7-Inputs!B19)',
                    lambda lv: max(0, lv['Caso8_PagoInferior']['AI7'].value - lv['Inputs']['B19'].value),
                    periodos_adicionales,
                    'el Excel no la calcula: último periodo con pago (AI7) − cuotas pactadas, mínimo 0')),
    29: (9, columna('Cobranza causada', 'cobranzaCausada')),
    30: (9, INTERES_MORA),
    31: (9, columna('Pago capital', 'pagoCapital')),
    32: (9, CAPITAL_VENCIDO),
    33: (9, celda('AL10', estado_final, 'estado al horizonte (120 semanas)')),
    34: (10, celda('U4', primer_cambio_de_tasa)),
    35: (10, CUOTA_FINANCIERA),
    36: (10, celda('U6', ultimo_periodo)),
    37: (10, SALDO_FINAL),
    38: (11, columna('Pago total cliente', 'pagoRecibido', 'la hoja del caso 11 no tiene "Pago recibido": es "Pago total cliente"')),
    39: (11, CAPITAL_VENCIDO),
    40: (11, CUOTAS_VENCIDAS),
    41: (11, DEUDA_A_EXTINGUIR),
    42: (11, CXC_FCC),
    43: (12, PAGO_RECIBIDO),
    44: (12, CAPITAL_VENCIDO),
    45: (12, columna('Saldo capital final', 'saldoFinal')),
}


def mapeo_de(numero: int, caso_excel: str) -> tuple[Optional[Mapeo], str]:
    """El mapeo de una prueba, o None con el motivo si no se puede mapear con certeza."""
    if numero not in MAPEO:
        return None, f'sin mapeo: la prueba {numero} no está en la tabla de la herramienta'
    caso, fabrica = MAPEO[numero]
    if caso_excel != f'Caso {caso}':
        return None, f'sin mapeo: la prueba {numero} es del "{caso_excel}" y la herramienta la tiene en el caso {caso}'
    return fabrica(HOJAS[caso]), ''


# ----------------------------------------------------------------------------------------------- comparación

def a_decimal(valor) -> Optional[Decimal]:
    if valor is None or valor == '':
        return None
    if isinstance(valor, float):
        return Decimal(repr(valor))
    return Decimal(str(valor))


def comparar(tipo: str, motor, esperado, tolerancia: Decimal) -> tuple[Optional[Decimal], str]:
    """(diferencia, 'OK' | 'FALLA'). En 'Exacto' compara texto normalizado; en el resto, |motor − esperado|."""
    if tipo == 'Exacto' and not isinstance(motor, Decimal):
        diferencia = Decimal(0) if normalizar_estado(motor) == normalizar_estado(esperado) else Decimal(1)
    else:
        diferencia = abs(a_decimal(motor) - a_decimal(esperado))
    return diferencia, ('OK' if diferencia <= tolerancia else 'FALLA')


@dataclass
class Resultado:
    numero: int
    fila: int
    caso: str
    variable: str
    tipo: str
    periodo: Optional[int]
    motor: object = None
    esperado: object = None
    diferencia: Optional[Decimal] = None
    tolerancia: Decimal = Decimal(0)
    estado: str = SIN_MAPEO
    nota: str = ''
    origen: str = ''
    formula_esperado: str = ''


def filas_de_pruebas(hoja):
    """Filas de la tabla principal: termina en la primera fila sin número en la columna A."""
    fila = PRIMERA_FILA
    while isinstance(hoja.cell(fila, COL['numero']).value, int):
        yield fila
        fila += 1


def leer_prueba(hoja, fila: int) -> Resultado:
    def valor(clave):
        return hoja.cell(fila, COL[clave]).value
    tipo, periodo, tolerancia = valor('tipo'), valor('periodo'), valor('tolerancia')
    return Resultado(numero=valor('numero'), fila=fila, caso=valor('caso'), variable=valor('variable'), tipo=tipo,
                     periodo=int(periodo) if isinstance(periodo, (int, float)) else None,
                     tolerancia=a_decimal(tolerancia) if tolerancia is not None
                     else TOLERANCIA_POR_TIPO.get(tipo, Decimal(0)))


def evaluar_prueba(r: Resultado, libro_valores, respuesta_del_caso: Callable[[int], dict]):
    """Llena el resultado de una prueba; si algo no se puede mapear, queda SIN MAPEO con el motivo."""
    mapeo, motivo = mapeo_de(r.numero, r.caso)
    if mapeo is None:
        r.nota = motivo
        return
    caso = int(r.caso.split()[-1])
    try:
        esperado = mapeo.esperado(libro_valores, r.periodo)
        r.formula_esperado = esperado.formula
        r.esperado = esperado.valor(libro_valores)
        r.motor = mapeo.motor(respuesta_del_caso(caso), r.periodo)
    except KeyError as e:
        r.nota = f'sin mapeo: {e.args[0]}'
        r.motor = None
        return
    r.diferencia, r.estado = comparar(r.tipo, r.motor, r.esperado, r.tolerancia)
    r.nota = mapeo.nota
    donde = f' (periodo {r.periodo})' if r.periodo is not None else ''
    r.origen = f'/casos/{caso}: {mapeo.motor.descripcion}{donde} ↔ {esperado.formula[1:]}'


def evaluar(libro_valores, obtener_caso: Callable[[int], dict]) -> list[Resultado]:
    """Calcula todas las pruebas. `obtener_caso(n)` devuelve la respuesta de /casos/{n} (se pide una vez)."""
    hoja = libro_valores[HOJA]
    respuestas = {}

    def respuesta_del_caso(caso):
        if caso not in respuestas:
            respuestas[caso] = obtener_caso(caso)
        return respuestas[caso]

    resultados = [leer_prueba(hoja, fila) for fila in filas_de_pruebas(hoja)]
    for r in resultados:
        evaluar_prueba(r, libro_valores, respuesta_del_caso)
    return resultados


# ----------------------------------------------------------------------------------------------- escritura

FORMATOS = {'Dinero': '#,##0.00', 'Tasa': '0.0000000000', 'Conteo': '0', 'Exacto': '@'}
VERDE = PatternFill('solid', start_color='C6EFCE', end_color='C6EFCE')
ROJO = PatternFill('solid', start_color='FFC7CE', end_color='FFC7CE')
GRIS = PatternFill('solid', start_color='EDEDED', end_color='EDEDED')


def escribir_encabezados(hoja):
    modelo = hoja.cell(FILA_ENCABEZADO, COL['observaciones'])
    for clave, texto in ENCABEZADOS_NUEVOS.items():
        c = hoja.cell(FILA_ENCABEZADO, COL[clave], texto)
        c.font, c.fill, c.border, c.alignment = copy(modelo.font), copy(modelo.fill), copy(modelo.border), copy(modelo.alignment)
    for clave, ancho in (('esperado', 18), ('diferencia', 14), ('estado', 14), ('origen', 70), ('observaciones', 45)):
        hoja.column_dimensions[get_column_letter(COL[clave])].width = ancho


def formula_diferencia(tipo: str, celda_motor: str, celda_excel: str) -> str:
    if tipo == 'Exacto':  # 0 si el texto coincide sin espacios ni tildes ("En mora" = EN_MORA)
        return (f'=IF({celda_motor}="","",IF(SUBSTITUTE(SUBSTITUTE(UPPER({celda_excel})," ","_"),"Í","I")'
                f'=SUBSTITUTE(UPPER({celda_motor})," ","_"),0,1))')
    return f'=IF(OR({celda_motor}="",{celda_excel}=""),"",ABS({celda_motor}-{celda_excel}))'


def escribir_fila(hoja, r: Resultado):
    f = r.fila
    F, G, K, L = (f'{get_column_letter(COL[c])}{f}' for c in ('resultado', 'tolerancia', 'esperado', 'diferencia'))
    if r.nota:
        hoja.cell(f, COL['observaciones'], r.nota)
    if r.estado == SIN_MAPEO:
        hoja[F] = None
        hoja.cell(f, COL['estado'], SIN_MAPEO)
        return
    formato = FORMATOS.get(r.tipo, 'General')
    if isinstance(r.motor, Decimal):
        hoja[F] = int(r.motor) if r.tipo == 'Conteo' else float(r.motor)
    else:
        hoja[F] = r.motor
    hoja[K] = r.formula_esperado
    hoja[L] = formula_diferencia(r.tipo, F, K)
    for ref in (F, K, L):
        hoja[ref].number_format = formato if r.tipo != 'Exacto' or ref != L else '0'
    hoja.cell(f, COL['estado'], f'=IF({L}="","",IF({L}<={G},"OK","FALLA"))')
    hoja.cell(f, COL['origen'], r.origen)


def escribir(hoja, resultados: list[Resultado]):
    escribir_encabezados(hoja)
    for r in resultados:
        escribir_fila(hoja, r)
    if resultados:
        letra = get_column_letter(COL['estado'])
        rango = f'{letra}{resultados[0].fila}:{letra}{resultados[-1].fila}'
        for texto, relleno in (('OK', VERDE), ('FALLA', ROJO), (SIN_MAPEO, GRIS)):
            hoja.conditional_formatting.add(rango, CellIsRule(operator='equal', formula=[f'"{texto}"'], fill=relleno))


def ruta_de_salida(excel: Path, salida: Path, perfil: str) -> Path:
    """Archivo nuevo en `salida`; si ya existe se numera. Nunca devuelve la ruta del original."""
    salida.mkdir(parents=True, exist_ok=True)
    base = f'{excel.stem}-resultados-{perfil.lower()}-{date.today().isoformat()}'
    n = 1
    destino = salida / f'{base}.xlsx'
    while destino.exists() or destino.resolve() == excel.resolve():
        n += 1
        destino = salida / f'{base}-{n}.xlsx'
    return destino


def guardar(excel: Path, destino: Path, resultados: list[Resultado]):
    if destino.resolve() == excel.resolve():
        raise ValueError('la salida no puede ser el Excel original')
    libro = openpyxl.load_workbook(excel)  # con fórmulas
    escribir(libro[HOJA], resultados)
    libro.calculation.fullCalcOnLoad = True  # openpyxl no guarda valores calculados: Excel recalcula al abrir
    libro.save(destino)


# ----------------------------------------------------------------------------------------------- API

def cliente_http(base: str, perfil: str, espera: int = 120) -> Callable[[int], dict]:
    def obtener(numero: int) -> dict:
        url = f'{base.rstrip("/")}/casos/{numero}?perfilReglas={perfil}'
        with urllib.request.urlopen(url, timeout=espera) as r:
            return json.loads(r.read().decode('utf-8'))
    return obtener


def resumen(resultados: list[Resultado]) -> str:
    ok = sum(r.estado == 'OK' for r in resultados)
    lineas = [f'{ok} de {len(resultados)} pruebas coinciden con el Excel']
    for r in resultados:
        if r.estado == 'OK':
            continue
        if r.estado == SIN_MAPEO:
            lineas.append(f'  #{r.numero} {r.caso} · {r.variable}: {r.nota}')
        else:
            lineas.append(f'  #{r.numero} {r.caso} · {r.variable}: FALLA motor={r.motor} excel={r.esperado} '
                          f'diferencia={r.diferencia} tolerancia={r.tolerancia}')
    return '\n'.join(lineas)


def main(argv=None):
    p = argparse.ArgumentParser(description='Llena "Resultado software" de la hoja Validaciones con el motor.')
    p.add_argument('excel', nargs='?', type=Path, default=EXCEL, help='Excel de Diego (por defecto docs/fuentes)')
    p.add_argument('--api', action='store_true', help=f'usar el API desplegado ({URL_AZURE})')
    p.add_argument('--url', default=None, help=f'base del API (por defecto {URL_LOCAL})')
    p.add_argument('--perfil', default='EXCEL', choices=['EXCEL', 'NEGOCIO'], help='perfil de reglas (EXCEL)')
    p.add_argument('--salida', type=Path, default=Path(tempfile.gettempdir()) / 'motor-cartera-validaciones',
                   help='directorio del archivo nuevo (por defecto el temporal del sistema)')
    a = p.parse_args(argv)
    sys.stdout.reconfigure(encoding='utf-8')
    base = a.url or (URL_AZURE if a.api else URL_LOCAL)
    print(f'API: {base} · perfil {a.perfil}')
    resultados = evaluar(openpyxl.load_workbook(a.excel, data_only=True), cliente_http(base, a.perfil))
    destino = ruta_de_salida(a.excel, a.salida, a.perfil)
    guardar(a.excel, destino, resultados)
    print(resumen(resultados))
    print(f'Excel con resultados: {destino}')
    return 0 if all(r.estado == 'OK' for r in resultados) else 1


if __name__ == '__main__':
    sys.exit(main())
