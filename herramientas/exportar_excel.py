"""Exporta el Excel de casos operativos a CSV para las pruebas del motor (HT-10).

Uso (desde la raíz del repo):
    python herramientas/exportar_excel.py [ruta_excel]

Por defecto lee docs/fuentes/casos-operaciones-v2.xlsx y escribe en
src/test/resources/excel/:
  - <Hoja>.csv por cada hoja de caso: la tabla del cronograma (fila 2 = encabezados,
    desde la fila 3), con los valores calculados por Excel.
  - inputs.csv: parámetros de la hoja Inputs (celda, etiqueta, valor).
  - validaciones.csv: la tabla de la hoja Validaciones.
Los números se escriben con repr() de Python para no perder precisión.
"""
import csv, sys
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
EXCEL = Path(sys.argv[1]) if len(sys.argv) > 1 else RAIZ / 'docs' / 'fuentes' / 'casos-operaciones-v2.xlsx'
SALIDA = RAIZ / 'src' / 'test' / 'resources' / 'excel'


def valor(v):
    if v is None:
        return ''
    if isinstance(v, float):
        return repr(v)
    return str(v)


def exportar_caso(hoja, destino):
    # La tabla termina en la primera columna vacía de la fila de encabezados.
    columnas = []
    c = 1
    while hoja.cell(2, c).value is not None:
        columnas.append(c)
        c += 1
    with destino.open('w', encoding='utf-8', newline='') as f:
        w = csv.writer(f)
        w.writerow([hoja.cell(2, c).value for c in columnas])
        for r in range(3, hoja.max_row + 1):
            if hoja.cell(r, 1).value is None:
                continue
            fila = [valor(hoja.cell(r, c).value) for c in columnas]
            # Filas sin datos (fuera del horizonte del caso) se omiten.
            if all(x in ('', '0') for x in fila[2:]):
                continue
            w.writerow(fila)


def exportar_inputs(hoja, destino):
    with destino.open('w', encoding='utf-8', newline='') as f:
        w = csv.writer(f)
        w.writerow(['celda', 'etiqueta', 'valor'])
        for fila in hoja.iter_rows(min_row=1, max_row=hoja.max_row, max_col=4):
            for c in fila[1:]:
                if isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                    etiqueta = fila[0].value if c.column == 2 else (f'{fila[0].value} · {fila[1].value}')
                    w.writerow([c.coordinate, etiqueta, valor(c.value)])


def exportar_validaciones(hoja, destino):
    with destino.open('w', encoding='utf-8', newline='') as f:
        w = csv.writer(f)
        for r in range(2, hoja.max_row + 1):
            w.writerow([valor(hoja.cell(r, c).value) for c in range(1, 11)])


def main():
    libro = openpyxl.load_workbook(EXCEL, data_only=True)
    SALIDA.mkdir(parents=True, exist_ok=True)
    for hoja in libro.worksheets:
        if hoja.title.startswith('Caso') and '_' in hoja.title:
            exportar_caso(hoja, SALIDA / f'{hoja.title}.csv')
            print('caso', hoja.title)
    exportar_inputs(libro['Inputs'], SALIDA / 'inputs.csv')
    exportar_validaciones(libro['Validaciones'], SALIDA / 'validaciones.csv')
    print('listo en', SALIDA)


if __name__ == '__main__':
    main()
